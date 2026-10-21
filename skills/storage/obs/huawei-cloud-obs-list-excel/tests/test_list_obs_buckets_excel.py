import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))

from list_obs_buckets_excel import (  # noqa: E402
    ALLOWED_OPERATION,
    ApiError,
    AuthError,
    ConfigError,
    ExportError,
    GuardError,
    ObsListError,
    _screen_output_value,
    build_command,
    classify_hcloud_failure,
    default_output_path,
    export_xlsx,
    guard_operation,
    guard_write_intent,
    list_buckets,
    list_buckets_via_sdk,
    main,
    normalize_sdk_bucket,
    parse_cli_buckets,
    render_table,
    run_hcloud,
)

CLI_OUTPUT = """Bucket number is: 2

  obs://bucket-alpha
  obs://bucket-beta
"""

CLI_OUTPUT_WITH_COLUMNS = """Bucket number is: 3

  obs://bucket-alpha  2024-01-02T03:04:05Z  cn-north-4  Standard
  obs://bucket-beta   2024-05-06T07:08:09Z  cn-south-1  Warm
  obs://bucket-gamma  2024-07-08T09:10:11Z  cn-east-3  Standard  private
"""


def hcloud_result(stdout="", stderr="", returncode=0):
    return subprocess.CompletedProcess(
        args=["hcloud"], returncode=returncode, stdout=stdout, stderr=stderr
    )


class FakeRunner:
    """subprocess.run fake returning canned hcloud outputs per call."""

    def __init__(self, results):
        self.results = list(results)
        self.calls = []

    def __call__(self, cmd, **kwargs):
        self.calls.append(cmd)
        if not self.results:
            raise AssertionError("unexpected extra hcloud call: %s" % cmd)
        out = self.results.pop(0)
        if isinstance(out, Exception):
            raise out
        return hcloud_result(*out)


class TestScreenOutputValue(unittest.TestCase):
    def test_drops_output_value(self):
        screened = _screen_output_value(
            ["--region", "cn-north-4", "--output", "/tmp/delete.xlsx"])
        self.assertNotIn("/tmp/delete.xlsx", screened)
        self.assertIn("--output", screened)

    def test_output_equals_form(self):
        screened = _screen_output_value(["--output=/tmp/downloads/x.xlsx", "--region", "cn-north-4"])
        self.assertIn("--output=<path>", screened)
        self.assertFalse(any("downloads" in t for t in screened))

    def test_short_o_value_dropped(self):
        screened = _screen_output_value(["-o", "/tmp/remove.xlsx"])
        self.assertNotIn("/tmp/remove.xlsx", screened)

    def test_other_args_kept(self):
        screened = _screen_output_value(["--page-size", "10", "--summary"])
        self.assertEqual(screened, ["--page-size", "10", "--summary"])


class TestWriteIntentGuard(unittest.TestCase):
    def test_read_intent_passes(self):
        self.assertFalse(guard_write_intent("--region cn-north-4 --format table"))
        self.assertFalse(guard_write_intent("list buckets show get describe"))

    def test_create_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("CreateBucket my-bucket")

    def test_delete_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("--operation DeleteBucket")

    def test_mutating_verbs_blocked(self):
        for text in ("upload object", "download object", "copy object",
                     "remove bucket", "restore object", "abort upload",
                     "sync dir", "modify acl", "update policy"):
            with self.assertRaises(GuardError):
                guard_write_intent(text)

    def test_normalized_separators_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("--create-bucket")
        with self.assertRaises(GuardError):
            guard_write_intent("update_bucket_policy")

    def test_operation_guard_rejects_non_ls(self):
        with self.assertRaises(GuardError):
            guard_operation("ListBuckets")
        with self.assertRaises(GuardError):
            guard_operation("mb")
        self.assertTrue(guard_operation(ALLOWED_OPERATION))

    def test_output_path_value_not_guarded(self):
        self.assertFalse(guard_write_intent(
            " ".join(_screen_output_value(
                ["--output", "/tmp/downloads/report.xlsx"]))))

    def test_region_keyword_still_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent(
                " ".join(_screen_output_value(["--region", "delete-me"])))


class TestBuildCommand(unittest.TestCase):
    def test_basic(self):
        argv = build_command("cn-north-4")
        self.assertEqual(argv[0], "OBS")
        self.assertEqual(argv[1], "ls")
        self.assertIn("-limit=1000", argv)
        self.assertIn("-e=obs.cn-north-4.myhuaweicloud.com", argv)

    def test_region_endpoint(self):
        argv = build_command("cn-south-1", limit=5)
        self.assertIn("-limit=5", argv)
        self.assertIn("-e=obs.cn-south-1.myhuaweicloud.com", argv)

    def test_custom_endpoint(self):
        argv = build_command("cn-north-4", endpoint="obs.custom.example.com")
        self.assertIn("-e=obs.custom.example.com", argv)

    def test_no_write_flags_ever_added(self):
        argv = build_command("cn-north-4")
        for flag in argv:
            self.assertFalse(
                flag.lower().lstrip("-").startswith(("create", "delete", "put", "rm")))


class TestRunHcloudEnvArgs(unittest.TestCase):
    def _fake_run(self):
        calls = []

        def fake(cmd, **kwargs):
            calls.append(cmd)
            return hcloud_result(stdout=CLI_OUTPUT)

        return calls, fake

    def test_extra_args_appended_after_readonly_argv(self):
        calls, fake = self._fake_run()
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake), \
                mock.patch.dict(os.environ,
                                {"OBS_LIST_EXCEL_HCLOUD_ARGS": "-v"},
                                clear=False):
            run_hcloud(build_command("cn-north-4"))
        self.assertEqual(calls[0][:3], ["hcloud", "OBS", "ls"])
        self.assertEqual(calls[0][-1], "-v")

    def test_extra_args_cannot_rewrite_service_or_operation(self):
        calls, fake = self._fake_run()
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake), \
                mock.patch.dict(os.environ,
                                {"OBS_LIST_EXCEL_HCLOUD_ARGS":
                                     "OBS cp obs://a obs://b"},
                                clear=False):
            run_hcloud(build_command("cn-north-4"))
        self.assertEqual(calls[0][1], "OBS")
        self.assertEqual(calls[0][2], "ls")
        self.assertNotIn("cp", calls[0][:3])

    def test_write_intent_in_extra_args_blocked(self):
        with mock.patch.dict(os.environ,
                             {"OBS_LIST_EXCEL_HCLOUD_ARGS":
                                  "obs://a delete"},
                             clear=False):
            with self.assertRaises(GuardError):
                run_hcloud(build_command("cn-north-4"))


class TestParseCli(unittest.TestCase):
    def test_names_only_output(self):
        buckets = parse_cli_buckets(CLI_OUTPUT)
        self.assertEqual([b["name"] for b in buckets], ["bucket-alpha", "bucket-beta"])

    def test_bucket_number_line_ignored(self):
        buckets = parse_cli_buckets("Bucket number is: 0\n")
        self.assertEqual(buckets, [])

    def test_columns_extracted(self):
        buckets = parse_cli_buckets(CLI_OUTPUT_WITH_COLUMNS)
        self.assertEqual(buckets[0]["name"], "bucket-alpha")
        self.assertEqual(buckets[0]["creation_date"], "2024-01-02T03:04:05Z")
        self.assertEqual(buckets[0]["location"], "cn-north-4")
        self.assertEqual(buckets[0]["storage_class"], "Standard")
        self.assertEqual(buckets[1]["storage_class"], "Warm")
        self.assertEqual(buckets[2]["policy"], "private")
        self.assertNotIn("policy", buckets[0])

    def test_duplicates_collapse(self):
        text = "  obs://dup\n  obs://dup\n"
        buckets = parse_cli_buckets(text)
        self.assertEqual(len(buckets), 1)

    def test_empty_output(self):
        self.assertEqual(parse_cli_buckets(""), [])


class TestClassifyFailure(unittest.TestCase):
    def test_missing_credentials(self):
        err = classify_hcloud_failure(
            "", "Warn: Please set ak, sk and endpoint in the configuration file!", 3)
        self.assertIsInstance(err, ConfigError)
        self.assertIn("AK/SK", str(err))

    def test_invalid_ak(self):
        err = classify_hcloud_failure("", "InvalidAccessKeyId: bad key", 1)
        self.assertIsInstance(err, AuthError)

    def test_auth_403(self):
        err = classify_hcloud_failure("", "AccessDenied 403 Forbidden", 1)
        self.assertIsInstance(err, AuthError)
        self.assertIn("IAM", str(err))

    def test_generic_api_error(self):
        err = classify_hcloud_failure("", "boom", 2)
        self.assertIsInstance(err, ApiError)

    def test_use_error_marker_with_zero_rc(self):
        err = classify_hcloud_failure(
            "[USE_ERROR]Invalid parameter: foo", "", 0)
        self.assertIsInstance(err, ApiError)
        self.assertIn("参数错误", str(err))


class TestListBuckets(unittest.TestCase):
    def test_success(self):
        fake = FakeRunner([(CLI_OUTPUT, "", 0)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            buckets = list_buckets("cn-north-4")
        self.assertEqual([b["name"] for b in buckets], ["bucket-alpha", "bucket-beta"])
        self.assertEqual(len(fake.calls), 1)
        self.assertIn("OBS", fake.calls[0])

    def test_empty(self):
        fake = FakeRunner([("Bucket number is: 0\n", "", 0)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            buckets = list_buckets("cn-north-4")
        self.assertEqual(buckets, [])

    def test_missing_credentials_raises_config(self):
        fake = FakeRunner([("", "Warn: Please set ak, sk and endpoint", 3)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            with self.assertRaises(ConfigError):
                list_buckets("cn-north-4")

    def test_auth_failure_raises_auth(self):
        fake = FakeRunner([("", "403 Forbidden", 1)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            with self.assertRaises(AuthError):
                list_buckets("cn-north-4")

    def test_use_error_rc0_raises_api(self):
        fake = FakeRunner([("[USE_ERROR]Invalid parameter: foo", "", 0)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            with self.assertRaises(ApiError):
                list_buckets("cn-north-4")

    def test_timeout(self):
        fake = FakeRunner([subprocess.TimeoutExpired(cmd="hcloud", timeout=120)])
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake):
            with self.assertRaises(ApiError):
                list_buckets("cn-north-4")

    def test_sdk_fallback_when_hcloud_missing(self):
        with mock.patch("list_obs_buckets_excel.subprocess.run",
                        FakeRunner([FileNotFoundError()])), \
                mock.patch("list_obs_buckets_excel.list_buckets_via_sdk",
                           return_value=[{"name": "sdk-bucket"}]) as sdk_mock:
            buckets = list_buckets("cn-north-4")
        self.assertEqual(buckets, [{"name": "sdk-bucket"}])
        sdk_mock.assert_called_once()

    def test_sdk_fallback_unavailable_raises_config(self):
        with mock.patch("list_obs_buckets_excel.subprocess.run",
                        FakeRunner([FileNotFoundError()])), \
                mock.patch("list_obs_buckets_excel.list_buckets_via_sdk",
                           side_effect=ConfigError("huaweicloudsdkobs 未安装")):
            with self.assertRaises(ConfigError):
                list_buckets("cn-north-4")


class TestSdkFallback(unittest.TestCase):
    def test_missing_env_credentials_raises_config(self):
        # Only blank the known AK/SK environment variable names instead of
        # enumerating os.environ (skillspector E2); empty values are treated
        # as missing by _env_credentials().
        credential_vars = (
            "HW_ACCESS_KEY", "HUAWEI_ACCESS_KEY", "HUAWEICLOUD_SDK_AK",
            "HUAWEI_CLOUD_ACCESS_KEY",
            "HW_SECRET_KEY", "HUAWEI_SECRET_KEY", "HUAWEICLOUD_SDK_SK",
            "HUAWEI_CLOUD_SECRET_KEY",
            "HW_SECURITY_TOKEN", "HUAWEI_SECURITY_TOKEN",
            "HUAWEICLOUD_SDK_SECURITY_TOKEN",
        )
        with mock.patch.dict(os.environ,
                             {name: "" for name in credential_vars}):
            with self.assertRaises(ConfigError) as ctx:
                list_buckets_via_sdk("cn-north-4")
        self.assertIn("AK/SK", str(ctx.exception))

    def test_normalize_sdk_bucket(self):
        class B:
            name = "b1"
            creation_date = "2024-01-01T00:00:00Z"
            location = "cn-north-4"
            storage_class = "Standard"
            policy = "private"
        record = normalize_sdk_bucket(B())
        self.assertEqual(record["name"], "b1")
        self.assertEqual(record["creation_date"], "2024-01-01T00:00:00Z")
        self.assertEqual(record["location"], "cn-north-4")
        self.assertEqual(record["storage_class"], "Standard")
        self.assertEqual(record["policy"], "private")

    def test_normalize_sdk_bucket_minimal(self):
        class B:
            pass
        self.assertEqual(normalize_sdk_bucket(B()), {"name": "-"})


class TestExportXlsx(unittest.TestCase):
    def _buckets(self):
        return [
            {"name": "bucket-alpha", "creation_date": "2024-01-02T03:04:05Z",
             "location": "cn-north-4", "storage_class": "Standard",
             "policy": "private"},
            {"name": "bucket-beta", "location": "cn-south-1",
             "storage_class": "Warm"},
        ]

    def test_export_writes_valid_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "obs.xlsx")
            written = export_xlsx(self._buckets(), path, summary=True, region="cn-north-4")
            self.assertTrue(os.path.isfile(path))
            self.assertEqual(written, os.path.abspath(path))

            from openpyxl import load_workbook
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["OBS 桶列表"])
            ws = wb.active
            headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
            self.assertEqual(headers, ["桶名称", "创建时间", "区域", "存储类别", "桶策略"])

            row2 = [ws.cell(row=2, column=c).value for c in range(1, 6)]
            self.assertEqual(row2[0], "bucket-alpha")
            self.assertEqual(row2[1], "2024-01-02T03:04:05Z")
            self.assertEqual(row2[2], "cn-north-4")
            self.assertEqual(row2[3], "Standard")
            self.assertEqual(row2[4], "private")

            row3 = [ws.cell(row=3, column=c).value for c in range(1, 6)]
            self.assertEqual(row3[0], "bucket-beta")
            self.assertEqual(row3[3], "Warm")
            self.assertEqual(row3[4], "-")

            # summary row (bolded footer)
            row4 = [ws.cell(row=4, column=c).value for c in range(1, 6)]
            self.assertEqual(row4[0], "桶数量")
            self.assertEqual(row4[1], "2")
            self.assertEqual(row4[2], "Region cn-north-4")

    def test_export_no_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "obs.xlsx")
            export_xlsx(self._buckets(), path, summary=False)
            from openpyxl import load_workbook
            ws = load_workbook(path).active
            self.assertEqual(ws.max_row, 3)
            self.assertNotIn("桶数量", [ws.cell(row=r, column=1).value
                                       for r in range(1, ws.max_row + 1)])

    def test_export_empty_list_headers_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "obs.xlsx")
            export_xlsx([], path)
            from openpyxl import load_workbook
            ws = load_workbook(path).active
            self.assertEqual(ws.max_row, 1)
            self.assertEqual(ws.cell(row=1, column=1).value, "桶名称")

    def test_export_creates_missing_dir(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "nested", "dir", "obs.xlsx")
            export_xlsx(self._buckets(), path)
            self.assertTrue(os.path.isfile(path))

    def test_export_bad_path_raises_and_leaves_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            blocker = os.path.join(tmp, "blocker")
            with open(blocker, "w") as fh:
                fh.write("x")
            bad = os.path.join(blocker, "sub", "obs.xlsx")
            with self.assertRaises(ExportError):
                export_xlsx(self._buckets(), bad)
            self.assertFalse(os.path.exists(bad))

    def test_export_missing_openpyxl_hint(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "obs.xlsx")
            with mock.patch("export_xlsx._require_openpyxl",
                            side_effect=ExportError(
                                "缺少依赖 openpyxl，无法生成 Excel 文件")):
                with self.assertRaises(ExportError) as ctx:
                    export_xlsx(self._buckets(), path)
            self.assertIn("openpyxl", str(ctx.exception))
            self.assertFalse(os.path.exists(path))

    def test_default_output_path_format(self):
        path = default_output_path("cn-north-4")
        self.assertTrue(re.match(r"obs_buckets_cn-north-4_\d{8}_\d{6}\.xlsx$", path))


class TestRenderTable(unittest.TestCase):
    def test_table_contains_fields(self):
        buckets = parse_cli_buckets(CLI_OUTPUT_WITH_COLUMNS)
        table = render_table(buckets)
        for fragment in ("name", "creation_date", "location",
                         "bucket-alpha", "2024-01-02T03:04:05Z", "cn-north-4"):
            self.assertIn(fragment, table)


class TestMain(unittest.TestCase):
    def _run(self, argv, results):
        fake = FakeRunner(results)
        with mock.patch("list_obs_buckets_excel.subprocess.run", fake), \
                mock.patch("sys.stdout") as stdout, \
                mock.patch("sys.stderr") as stderr:
            rc = main(argv)
        out = "".join(c[0][0] for c in stdout.write.call_args_list)
        err = "".join(c[0][0] for c in stderr.write.call_args_list)
        return rc, out, err

    def test_xlsx_output_writes_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "obs.xlsx")
            rc, stdout, err = self._run(
                ["--region", "cn-north-4", "--output", out, "--summary"],
                [(CLI_OUTPUT_WITH_COLUMNS, "", 0)])
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.isfile(out))
            self.assertIn("obs.xlsx", stdout)
            self.assertIn("3 个桶", stdout)
            self.assertEqual(err, "")

    def test_xlsx_default_path_and_table_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = os.getcwd()
            os.chdir(tmp)
            try:
                import list_obs_buckets_excel as mod
                with mock.patch.object(mod, "default_output_path",
                                       return_value=os.path.join(tmp, "default.xlsx")):
                    fake = FakeRunner([(CLI_OUTPUT, "", 0)])
                    with mock.patch("list_obs_buckets_excel.subprocess.run", fake), \
                            mock.patch("sys.stdout") as stdout:
                        rc = mod.main([])
                    out = "".join(c[0][0] for c in stdout.write.call_args_list)
                    self.assertEqual(rc, 0)
                    self.assertIn("default.xlsx", out)
                    self.assertTrue(os.path.isfile(os.path.join(tmp, "default.xlsx")))
            finally:
                os.chdir(old)

        # table mode
        rc, out, err = self._run(["--format", "table"], [(CLI_OUTPUT, "", 0)])
        self.assertEqual(rc, 0)
        self.assertIn("OBS 桶列表", out)
        self.assertIn("bucket-alpha", out)

    def test_json_output(self):
        rc, out, err = self._run(["--format", "json"], [(CLI_OUTPUT, "", 0)])
        self.assertEqual(rc, 0)
        payload = json.loads(out)
        self.assertEqual([b["name"] for b in payload], ["bucket-alpha", "bucket-beta"])

    def test_empty_xlsx_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "empty.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [("Bucket number is: 0\n", "", 0)])
            self.assertEqual(rc, 0)
            self.assertIn("无 OBS 桶", stdout)
            self.assertFalse(os.path.exists(out))

    def test_auth_error_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "nope.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [("", "403 Forbidden", 1)])
            self.assertEqual(rc, 1)
            self.assertIn("IAM", err)
            self.assertFalse(os.path.exists(out))

    def test_missing_ak_sk_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "nope.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [("", "Please set ak, sk and endpoint", 3)])
            self.assertEqual(rc, 1)
            self.assertIn("AK/SK", err)
            self.assertFalse(os.path.exists(out))

    def test_hcloud_not_installed(self):
        with mock.patch("list_obs_buckets_excel.list_buckets_via_sdk",
                        side_effect=ConfigError("huaweicloudsdkobs 未安装")):
            rc, out, err = self._run([], [FileNotFoundError()])
        self.assertEqual(rc, 1)
        self.assertIn("未安装", err)

    def test_hcloud_missing_falls_back_to_sdk(self):
        with mock.patch("list_obs_buckets_excel.list_buckets_via_sdk",
                        return_value=[{"name": "sdk-bucket"}]):
            rc, out, err = self._run(["--format", "table"], [FileNotFoundError()])
        self.assertEqual(rc, 0)
        self.assertIn("sdk-bucket", out)

    def test_write_intent_operation_blocked(self):
        rc, out, err = self._run(["--operation", "CreateBucket"], [])
        self.assertEqual(rc, 1)
        self.assertIn("写操作意图", err)

    def test_write_keyword_argv_blocked(self):
        rc, out, err = self._run(["--region", "delete-me"], [])
        self.assertEqual(rc, 1)
        self.assertIn("守卫拦截", err)

    def test_output_path_with_keyword_not_blocked(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "downloads", "report.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [(CLI_OUTPUT, "", 0)])
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.isfile(out))


class TestErrorHierarchy(unittest.TestCase):
    def test_subclasses(self):
        for exc in (GuardError, ConfigError, AuthError, ApiError):
            self.assertTrue(issubclass(exc, ObsListError))


if __name__ == "__main__":
    unittest.main()