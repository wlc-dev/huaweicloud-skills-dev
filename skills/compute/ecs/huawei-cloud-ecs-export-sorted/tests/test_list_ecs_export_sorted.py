import csv
import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

# Never emit __pycache__/*.pyc for the bundled scripts modules.
sys.dont_write_bytecode = True

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))

import export_report  # noqa: E402
from export_report import (  # noqa: E402
    ExportError,
    REPORT_HEADERS,
    default_output_path,
    export_csv,
    export_xlsx,
)
from list_ecs_export_sorted import (  # noqa: E402
    ALLOWED_OPERATION,
    ApiError,
    AuthError,
    ConfigError,
    EcsExportError,
    GuardError,
    SortError,
    active_filter_descriptions,
    build_command,
    classify_hcloud_failure,
    extract_json,
    extract_servers,
    guard_operation,
    guard_text_from_argv,
    guard_write_intent,
    list_servers,
    list_servers_via_sdk,
    main,
    normalize_server,
    render_table,
    resolve_sort_field,
    run_hcloud,
    sdk_server_to_dict,
    sort_servers,
)

SERVER_1 = {
    "id": "ecs-0001",
    "name": "ecs-web-01",
    "status": "ACTIVE",
    "flavor": {"id": "s6.large.2", "name": "s6.large.2", "vcpus": "2", "ram": "4096"},
    "addresses": {
        "vpc-subnet": [
            {"addr": "192.168.0.10", "OS-EXT-IPS:type": "fixed", "version": 4}
        ]
    },
    "OS-EXT-AZ:availability_zone": "cn-north-4a",
    "created": "2023-03-01T10:00:00Z",
    "image": {"id": "img-0001"},
    "key_name": "kp-demo",
    "enterprise_project_id": "0",
}
SERVER_2 = {
    "id": "ecs-0002",
    "name": "ecs-db-02",
    "status": "SHUTOFF",
    "flavor": {"id": "s6.xlarge.2", "name": "s6.xlarge.2", "vcpus": "4", "ram": "8192"},
    "addresses": {
        "vpc-subnet": [
            {"addr": "192.168.0.11", "OS-EXT-IPS:type": "fixed"},
            {"addr": "10.0.0.5", "OS-EXT-IPS:type": "floating"},
        ]
    },
    "OS-EXT-AZ:availability_zone": "cn-north-4b",
    "created": "2023-01-01T10:00:00Z",
    "tags": ["env:test"],
}


def hcloud_result(stdout="", stderr="", returncode=0):
    return subprocess.CompletedProcess(
        args=["hcloud"], returncode=returncode, stdout=stdout, stderr=stderr
    )


class FakeRunner:
    """subprocess.run fake returning canned hcloud outputs per call."""

    def __init__(self, results):
        self.results = list(results)
        self.calls = []

    def __call__(self, cmd, **kwargs):
        self.calls.append(cmd)
        if not self.results:
            raise AssertionError("unexpected extra hcloud call: %s" % cmd)
        out = self.results.pop(0)
        if isinstance(out, Exception):
            raise out
        return hcloud_result(*out)


class TestWriteIntentGuard(unittest.TestCase):
    def test_read_intent_passes(self):
        self.assertFalse(guard_write_intent("--region cn-north-4 --format table list servers"))
        self.assertFalse(guard_write_intent("show server details flavors"))

    def test_create_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("CreateServers ecs-new")

    def test_delete_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("--operation DeleteServers")

    def test_mutating_verbs_blocked(self):
        for text in ("resize flavor", "reboot ecs-01", "stop server",
                     "update server name", "attach volume", "server-action"):
            with self.assertRaises(GuardError):
                guard_write_intent(text)

    def test_normalized_separators_blocked(self):
        with self.assertRaises(GuardError):
            guard_write_intent("--create-server")
        with self.assertRaises(GuardError):
            guard_write_intent("delete_server")

    def test_operation_guard_rejects_non_list(self):
        with self.assertRaises(GuardError):
            guard_operation("ShowServer")
        with self.assertRaises(GuardError):
            guard_operation("BatchStartServers")
        self.assertTrue(guard_operation(ALLOWED_OPERATION))

    def test_guard_text_strips_filter_values(self):
        self.assertEqual(guard_text_from_argv(["--status", "REBUILD"]), "")
        self.assertEqual(guard_text_from_argv(["--status=RESIZE"]), "")
        self.assertEqual(guard_text_from_argv(["--name", "restart-me"]), "")
        self.assertEqual(guard_text_from_argv(["--region", "cn-north-4"]), "")
        self.assertEqual(guard_text_from_argv(["--sort-by", "created"]), "")
        self.assertEqual(guard_text_from_argv(["--order", "desc"]), "")
        # Operation name is still scanned (defense in depth).
        self.assertIn("DeleteServers",
                      guard_text_from_argv(["--operation", "DeleteServers"]))
        # Flag names are scanned.
        self.assertIn("--create-server",
                      guard_text_from_argv(["--create-server"]))

    def test_status_rebuild_not_blocked(self):
        self.assertFalse(
            guard_write_intent(guard_text_from_argv(["--status", "REBUILD"])))
        self.assertFalse(
            guard_write_intent(guard_text_from_argv(["--name", "stop-me"])))

    def test_guard_text_strips_output_values(self):
        self.assertEqual(
            guard_text_from_argv(["--output", "ecs-delete-report.xlsx"]), "")
        self.assertEqual(
            guard_text_from_argv(["--output=ecs-start-stop-report.xlsx"]), "")
        self.assertEqual(
            guard_text_from_argv(["-o", "ecs-reboot-report.xlsx"]), "")
        # A following *flag* is still scanned: a mutating flag cannot hide
        # behind a value-stripped option.
        self.assertIn("--create-server",
                      guard_text_from_argv(["--output", "--create-server"]))
        self.assertIn("DeleteServers",
                      guard_text_from_argv(["--output", "--operation",
                                            "DeleteServers"]))

    def test_output_path_with_write_words_not_blocked(self):
        self.assertFalse(guard_write_intent(
            guard_text_from_argv(["--output", "ecs-delete-report.xlsx"])))
        self.assertFalse(guard_write_intent(
            guard_text_from_argv(["--output", "ecs-restart-report.xlsx"])))


class TestBuildCommand(unittest.TestCase):
    def test_basic(self):
        argv = build_command("cn-north-4")
        self.assertEqual(argv[0], "ECS")
        self.assertEqual(argv[1], "ListServersDetails")
        self.assertIn("--cli-region=cn-north-4", argv)
        self.assertIn("--limit=100", argv)
        self.assertIn("--offset=0", argv)

    def test_filters_appended(self):
        argv = build_command("cn-south-1", status="ACTIVE", name="ecs-web",
                             limit=10, offset=20)
        self.assertIn("--status=ACTIVE", argv)
        self.assertIn("--name=ecs-web", argv)
        self.assertIn("--limit=10", argv)
        self.assertIn("--offset=20", argv)

    def test_no_write_flags_ever_added(self):
        argv = build_command("cn-north-4", status="ACTIVE", name="ecs-web")
        for flag in argv:
            self.assertFalse(
                flag.lower().lstrip("-").startswith(
                    ("create", "delete", "update", "reboot", "resize")))


class TestRunHcloudEnvArgs(unittest.TestCase):
    def _fake_run(self):
        calls = []

        def fake(cmd, **kwargs):
            calls.append(cmd)
            return hcloud_result(stdout=json.dumps({"servers": [SERVER_1], "count": 1}))

        return calls, fake

    def test_extra_args_appended_after_readonly_argv(self):
        calls, fake = self._fake_run()
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake), \
                mock.patch.dict(os.environ,
                                {"ECS_EXPORT_SORTED_HCLOUD_ARGS": "-v"},
                                clear=False):
            run_hcloud(build_command("cn-north-4"))
        self.assertEqual(calls[0][:3], ["hcloud", "ECS", "ListServersDetails"])
        self.assertEqual(calls[0][-1], "-v")

    def test_extra_args_are_tail_and_cannot_rewrite_service_or_operation(self):
        calls, fake = self._fake_run()
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake), \
                mock.patch.dict(os.environ,
                                {"ECS_EXPORT_SORTED_HCLOUD_ARGS":
                                     "--cli-region=cn-south-1"},
                                clear=False):
            run_hcloud(build_command("cn-north-4"))
        self.assertEqual(calls[0][1], "ECS")
        self.assertEqual(calls[0][2], "ListServersDetails")
        self.assertEqual(calls[0][-1], "--cli-region=cn-south-1")

    def test_write_intent_in_extra_args_blocked(self):
        with mock.patch.dict(os.environ,
                             {"ECS_EXPORT_SORTED_HCLOUD_ARGS":
                                  "--operation DeleteServers"},
                             clear=False):
            with self.assertRaises(GuardError):
                run_hcloud(build_command("cn-north-4"))


class TestExtract(unittest.TestCase):
    def test_extract_json_plain(self):
        data = extract_json(json.dumps({"servers": []}))
        self.assertEqual(data, {"servers": []})

    def test_extract_json_with_noise(self):
        data = extract_json("WARNING noise\n{\"count\": 1}\n")
        self.assertEqual(data, {"count": 1})

    def test_extract_json_invalid(self):
        with self.assertRaises(ApiError):
            extract_json("<html>not json</html>")

    def test_extract_servers_known_keys(self):
        for key in ("servers", "server_list"):
            self.assertEqual(extract_servers({key: [SERVER_1]}), [SERVER_1])
        self.assertEqual(extract_servers([SERVER_1]), [SERVER_1])
        self.assertEqual(extract_servers({"count": 0}), [])


class TestClassifyFailure(unittest.TestCase):
    def test_missing_credentials(self):
        err = classify_hcloud_failure(
            "", "ERROR: cli-ak is empty, please run hcloud configure", 1)
        self.assertIsInstance(err, ConfigError)
        self.assertIn("AK/SK", str(err))

    def test_auth_403(self):
        err = classify_hcloud_failure(
            "", '{"error_code":"ApiGW.0301","error_msg":"403 Forbidden"}', 1)
        self.assertIsInstance(err, AuthError)
        self.assertIn("IAM", str(err))

    def test_auth_unauthorized(self):
        err = classify_hcloud_failure("", "Unauthorized: invalid AK", 1)
        self.assertIsInstance(err, AuthError)

    def test_generic_api_error(self):
        err = classify_hcloud_failure("", "boom", 2)
        self.assertIsInstance(err, ApiError)

    def test_use_error_marker_with_zero_rc(self):
        err = classify_hcloud_failure(
            "[USE_ERROR]Invalid parameter: xxx\n\n"
            "Run `hcloud ECS ListServersDetails --help` for details.", "", 0)
        self.assertIsInstance(err, ApiError)
        self.assertIn("参数错误", str(err))
        self.assertIn("--help", str(err))


class TestListServers(unittest.TestCase):
    def test_single_page(self):
        fake = FakeRunner([(
            json.dumps({"servers": [SERVER_1, SERVER_2], "count": 2}), "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            servers = list_servers("cn-north-4")
        self.assertEqual(servers, [SERVER_1, SERVER_2])
        self.assertEqual(len(fake.calls), 1)

    def test_pagination_via_count(self):
        page1 = {"servers": [SERVER_1], "count": 2}
        page2 = {"servers": [SERVER_2], "count": 2}
        fake = FakeRunner([(json.dumps(page1), "", 0), (json.dumps(page2), "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            servers = list_servers("cn-north-4", limit=1)
        self.assertEqual(servers, [SERVER_1, SERVER_2])
        self.assertEqual(len(fake.calls), 2)
        self.assertIn("--offset=1", fake.calls[1])

    def test_pagination_via_full_page_without_count(self):
        page1 = {"servers": [SERVER_1]}
        page2 = {"servers": []}
        fake = FakeRunner([(json.dumps(page1), "", 0), (json.dumps(page2), "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            servers = list_servers("cn-north-4", limit=1)
        self.assertEqual(servers, [SERVER_1])

    def test_empty_region(self):
        fake = FakeRunner([(json.dumps({"servers": [], "count": 0}), "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            servers = list_servers("cn-north-4")
        self.assertEqual(servers, [])

    def test_status_and_name_forwarded_server_side(self):
        fake = FakeRunner([(
            json.dumps({"servers": [SERVER_1], "count": 1}), "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            servers = list_servers("cn-north-4", status="ACTIVE", name="ecs-web")
        self.assertEqual(servers, [SERVER_1])
        call = fake.calls[0]
        self.assertIn("--status=ACTIVE", call)
        self.assertIn("--name=ecs-web", call)

    def test_use_error_rc0_raises_api_error(self):
        fake = FakeRunner([(
            "[USE_ERROR]Invalid parameter: xxx\n\n"
            "Run `hcloud ECS ListServersDetails --help` for details.", "", 0)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            with self.assertRaises(ApiError) as ctx:
                list_servers("cn-north-4")
        self.assertIn("参数错误", str(ctx.exception))

    def test_nonzero_exit_403(self):
        fake = FakeRunner([("", "403 Forbidden", 1)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            with self.assertRaises(AuthError):
                list_servers("cn-north-4")

    def test_nonzero_exit_credentials(self):
        fake = FakeRunner([("", "ERROR: no AK configured", 1)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            with self.assertRaises(ConfigError):
                list_servers("cn-north-4")

    def test_timeout(self):
        fake = FakeRunner([subprocess.TimeoutExpired(cmd="hcloud", timeout=120)])
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake):
            with self.assertRaises(ApiError):
                list_servers("cn-north-4")

    def test_sdk_fallback_when_hcloud_missing(self):
        with mock.patch("list_ecs_export_sorted.subprocess.run",
                        FakeRunner([FileNotFoundError()])), \
                mock.patch("list_ecs_export_sorted.list_servers_via_sdk",
                           return_value=[SERVER_1, SERVER_2]) as sdk_mock:
            servers = list_servers("cn-north-4", status="ACTIVE", name="ecs-web")
        self.assertEqual(servers, [SERVER_1, SERVER_2])
        self.assertEqual(sdk_mock.call_args.kwargs.get("status"), "ACTIVE")
        self.assertEqual(sdk_mock.call_args.kwargs.get("name"), "ecs-web")

    def test_sdk_fallback_unavailable_raises_config(self):
        with mock.patch("list_ecs_export_sorted.subprocess.run",
                        FakeRunner([FileNotFoundError()])), \
                mock.patch("list_ecs_export_sorted.list_servers_via_sdk",
                           side_effect=ConfigError("huaweicloudsdkecs 未安装")):
            with self.assertRaises(ConfigError):
                list_servers("cn-north-4")


class TestSdkFallback(unittest.TestCase):
    def test_missing_env_credentials_raises_config(self):
        credential_vars = (
            "HUAWEI_ACCESS_KEY", "HW_ACCESS_KEY", "HWC_ACCESS_KEY",
            "HUAWEICLOUD_SDK_AK", "HUAWEI_CLOUD_ACCESS_KEY",
            "HUAWEI_SECRET_KEY", "HW_SECRET_KEY", "HWC_SECRET_KEY",
            "HUAWEICLOUD_SDK_SK", "HUAWEI_CLOUD_SECRET_KEY",
        )
        with mock.patch.dict(os.environ,
                             {name: "" for name in credential_vars}):
            with self.assertRaises(ConfigError) as ctx:
                list_servers_via_sdk("cn-north-4")
        self.assertIn("AK/SK", str(ctx.exception))


class _FakeAddress:
    """Duck-typed stand-in for huaweicloudsdkecs ServerAddress."""

    def __init__(self, addr, ip_type="fixed", version=4):
        self.addr = addr
        self.version = version
        self.os_ext_ip_stype = ip_type

    def to_dict(self):
        return {"addr": self.addr, "version": self.version,
                "os_ext_ip_stype": self.os_ext_ip_stype}


class _FakeServerDetail:
    """Duck-typed stand-in for huaweicloudsdkecs ServerDetail."""

    def __init__(self, data):
        self._data = data

    def to_dict(self):
        return dict(self._data)


class TestSdkNormalize(unittest.TestCase):
    def test_wire_name_dict_passes_through(self):
        flat = sdk_server_to_dict(dict(SERVER_1))
        self.assertEqual(flat["name"], "ecs-web-01")
        self.assertEqual(flat["OS-EXT-AZ:availability_zone"], "cn-north-4a")

    def test_nested_sdk_objects_are_deep_converted(self):
        # Regression: ServerDetail.to_dict() leaves nested ServerAddress model
        # objects in ``addresses`` → json.dumps / render_table raised
        # "TypeError: Object of type ServerAddress is not JSON serializable".
        server = _FakeServerDetail({
            "id": "ecs-sdk-1",
            "name": "ecs-sdk-01",
            "addresses": {
                "vpc-subnet": [
                    _FakeAddress("192.168.1.10", "fixed"),
                    _FakeAddress("10.1.1.5", "floating"),
                ]
            },
        })
        flat = sdk_server_to_dict(server)
        # Must be JSON serializable end to end.
        json.dumps(flat)
        self.assertEqual(flat["id"], "ecs-sdk-1")
        self.assertIsInstance(flat["addresses"]["vpc-subnet"][0], dict)
        self.assertIsInstance(flat["addresses"]["vpc-subnet"][1], dict)
        # The private/public IP extraction now works on the SDK path.
        normalized = normalize_server(flat)
        self.assertEqual(normalized["privateIps"], "192.168.1.10")
        self.assertEqual(normalized["publicIps"], "10.1.1.5")

    def test_nested_sdk_objects_render_table(self):
        server = _FakeServerDetail({
            "id": "ecs-sdk-2",
            "name": "ecs-sdk-02",
            "addresses": {
                "vpc-subnet": [_FakeAddress("192.168.2.20", "fixed")]
            },
        })
        table = render_table([sdk_server_to_dict(server)])
        self.assertIn("ecs-sdk-02", table)
        self.assertIn("192.168.2.20", table)

    def test_unknown_non_native_object_is_stringified(self):
        from list_ecs_export_sorted import _to_jsonable

        class Weird:
            pass

        value = _to_jsonable({"k": Weird(), "n": [Weird()]})
        json.dumps(value)
        self.assertIsInstance(value["k"], str)
        self.assertIsInstance(value["n"][0], str)

    def test_normalize_accepts_sdk_model_object(self):
        # Regression: a ServerDetail model handed straight to normalize_server
        # (not pre-converted by sdk_server_to_dict) crashed with
        # "TypeError: 'ServerDetail' object is not iterable" / "not JSON
        # serializable". normalize_server must deep-convert its own input.
        server = _FakeServerDetail({
            "id": "ecs-sdk-3",
            "name": "ecs-sdk-03",
            "addresses": {
                "vpc-subnet": [
                    _FakeAddress("192.168.3.30", "fixed"),
                    _FakeAddress("10.3.3.5", "floating"),
                ]
            },
        })
        flat = normalize_server(server)
        json.dumps(flat)  # must not raise
        self.assertEqual(flat["name"], "ecs-sdk-03")
        self.assertEqual(flat["privateIps"], "192.168.3.30")
        self.assertEqual(flat["publicIps"], "10.3.3.5")

    def test_sdk_model_render_and_sort(self):
        # render_table / sort_servers must also tolerate a raw SDK model object.
        server = _FakeServerDetail({
            "id": "ecs-sdk-4",
            "name": "ecs-sdk-04",
            "addresses": {"vpc-subnet": [_FakeAddress("192.168.4.40")]},
        })
        self.assertIn("ecs-sdk-04", render_table([server]))
        self.assertEqual(sort_servers([server], "name", "asc")[0]["name"],
                         "ecs-sdk-04")


class TestNormalizeAndRender(unittest.TestCase):
    def test_normalize_flattens_flavor_az_and_ips(self):
        flat = normalize_server(SERVER_1)
        self.assertEqual(flat["flavor"], "s6.large.2")
        self.assertEqual(flat["vcpus"], "2")
        self.assertEqual(flat["ram"], "4096")
        self.assertEqual(flat["availability_zone"], "cn-north-4a")
        self.assertEqual(flat["privateIps"], "192.168.0.10")
        self.assertEqual(flat["publicIps"], "-")
        self.assertEqual(flat["image_id"], "img-0001")
        self.assertEqual(flat["created"], "2023-03-01T10:00:00Z")

    def test_normalize_public_and_private_ips(self):
        flat = normalize_server(SERVER_2)
        self.assertEqual(flat["privateIps"], "192.168.0.11")
        self.assertEqual(flat["publicIps"], "10.0.0.5")

    def test_normalize_missing_fields_dash(self):
        flat = normalize_server({"id": "ecs-x"})
        self.assertEqual(flat["availability_zone"], "-")
        self.assertEqual(flat["privateIps"], "-")
        self.assertEqual(flat["publicIps"], "-")
        self.assertEqual(flat["created"], "-")

    def test_normalize_sdk_attr_names(self):
        flat = normalize_server({
            "id": "ecs-y",
            "os_ext_a_zavailability_zone": "cn-north-4c",
            "access_i_pv4": "9.9.9.9",
        })
        self.assertEqual(flat["availability_zone"], "cn-north-4c")
        self.assertEqual(flat["publicIps"], "9.9.9.9")

    def test_table_contains_basic_info(self):
        table = render_table([SERVER_1, SERVER_2])
        for fragment in ("ecs-web-01", "ecs-0001", "ACTIVE", "s6.large.2",
                         "cn-north-4a", "192.168.0.10", "ecs-db-02",
                         "SHUTOFF", "10.0.0.5"):
            self.assertIn(fragment, table)

    def test_active_filter_descriptions(self):
        self.assertEqual(active_filter_descriptions("ACTIVE", "ecs-web"),
                         ["status=ACTIVE", "name=ecs-web"])
        self.assertEqual(active_filter_descriptions(), [])
        self.assertEqual(
            active_filter_descriptions(sort_by="created", order="desc"),
            ["sort-by=created order=desc"])


class TestSorting(unittest.TestCase):
    def _records(self):
        return [
            normalize_server(SERVER_1),  # name ecs-web-01, vcpus 2, created 03-01
            normalize_server(SERVER_2),  # name ecs-db-02, vcpus 4, created 01-01
        ]

    def test_resolve_aliases(self):
        self.assertEqual(resolve_sort_field("name"), "name")
        self.assertEqual(resolve_sort_field("vCPU"), "vcpus")
        self.assertEqual(resolve_sort_field("创建时间"), "created")
        self.assertEqual(resolve_sort_field("可用区"), "availability_zone")
        self.assertIsNone(resolve_sort_field(None))

    def test_unknown_field_raises(self):
        with self.assertRaises(SortError) as ctx:
            resolve_sort_field("nope")
        self.assertIn("未知的排序字段", str(ctx.exception))

    def test_no_sort_preserves_order(self):
        servers = self._records()
        self.assertEqual(sort_servers(servers, None, "asc"), servers)

    def test_sort_text_asc_desc(self):
        servers = self._records()
        asc = sort_servers(servers, "name", "asc")
        self.assertEqual([s["name"] for s in asc], ["ecs-db-02", "ecs-web-01"])
        desc = sort_servers(servers, "name", "desc")
        self.assertEqual([s["name"] for s in desc], ["ecs-web-01", "ecs-db-02"])

    def test_sort_numeric_asc_desc(self):
        servers = self._records()
        asc = sort_servers(servers, "vcpus", "asc")
        self.assertEqual([s["name"] for s in asc], ["ecs-web-01", "ecs-db-02"])
        desc = sort_servers(servers, "vcpus", "desc")
        self.assertEqual([s["name"] for s in desc], ["ecs-db-02", "ecs-web-01"])

    def test_sort_datetime_asc_desc(self):
        servers = self._records()
        asc = sort_servers(servers, "created", "asc")
        self.assertEqual([s["name"] for s in asc], ["ecs-db-02", "ecs-web-01"])
        desc = sort_servers(servers, "created", "desc")
        self.assertEqual([s["name"] for s in desc], ["ecs-web-01", "ecs-db-02"])

    def test_invalid_values_always_last(self):
        servers = [
            {"name": "-"},
            {"name": "ecs-a"},
            {"name": None},
            {"name": "ecs-b"},
        ]
        for order in ("asc", "desc"):
            ordered = sort_servers(servers, "name", order)
            self.assertEqual([s["name"] for s in ordered[:2]],
                             ["ecs-a", "ecs-b"] if order == "asc" else ["ecs-b", "ecs-a"])
            self.assertEqual(ordered[2]["name"], "-")
            self.assertIsNone(ordered[3]["name"])

    def test_numeric_unparseable_last(self):
        servers = [
            {"vcpus": "abc"},
            {"vcpus": "4"},
            {"vcpus": "2"},
        ]
        ordered = sort_servers(servers, "vcpus", "asc")
        self.assertEqual([s["vcpus"] for s in ordered], ["2", "4", "abc"])

    def test_stable_sort_keeps_original_order_on_ties(self):
        servers = [
            {"name": "same", "id": "a"},
            {"name": "same", "id": "b"},
            {"name": "same", "id": "c"},
        ]
        ordered = sort_servers(servers, "name", "desc")
        self.assertEqual([s["id"] for s in ordered], ["a", "b", "c"])

    def test_invalid_order_raises(self):
        with self.assertRaises(SortError):
            sort_servers(self._records(), "name", "sideways")


class TestExportReport(unittest.TestCase):
    def _servers(self):
        return [
            normalize_server(SERVER_1),
            normalize_server(SERVER_2),
        ]

    def test_export_writes_valid_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ecs.xlsx")
            written = export_xlsx(self._servers(), path, summary=True,
                                  region="cn-north-4")
            self.assertTrue(os.path.isfile(path))
            self.assertEqual(written, os.path.abspath(path))

            from openpyxl import load_workbook
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["ECS 实例清单"])
            ws = wb.active
            headers = [ws.cell(row=1, column=c).value
                       for c in range(1, len(REPORT_HEADERS) + 1)]
            self.assertEqual(headers, REPORT_HEADERS)
            row2 = [ws.cell(row=2, column=c).value
                    for c in range(1, len(REPORT_HEADERS) + 1)]
            self.assertEqual(row2[0], "ecs-web-01")
            self.assertEqual(row2[3], "s6.large.2")
            self.assertEqual(row2[9], "2023-03-01T10:00:00Z")
            row4 = [ws.cell(row=4, column=c).value for c in range(1, 3)]
            self.assertEqual(row4, ["实例数量", "2"])

    def test_export_no_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ecs.xlsx")
            export_xlsx(self._servers(), path, summary=False)
            from openpyxl import load_workbook
            ws = load_workbook(path).active
            self.assertEqual(ws.max_row, 3)

    def test_export_creates_missing_dir(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "nested", "dir", "ecs.xlsx")
            export_xlsx(self._servers(), path)
            self.assertTrue(os.path.isfile(path))

    def test_export_bad_path_raises_and_leaves_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            blocker = os.path.join(tmp, "blocker")
            with open(blocker, "w") as fh:
                fh.write("x")
            bad = os.path.join(blocker, "sub", "ecs.xlsx")
            with self.assertRaises(ExportError):
                export_xlsx(self._servers(), bad)
            self.assertFalse(os.path.exists(bad))

    def test_export_missing_openpyxl_hint(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ecs.xlsx")
            with mock.patch("export_report._require_openpyxl",
                            side_effect=ExportError(
                                "缺少依赖 openpyxl，无法生成 Excel 文件")):
                with self.assertRaises(ExportError) as ctx:
                    export_xlsx(self._servers(), path)
            self.assertIn("openpyxl", str(ctx.exception))
            self.assertFalse(os.path.exists(path))

    def test_default_output_path_format(self):
        self.assertTrue(re.match(r"ecs_instances_cn-north-4_\d{8}_\d{6}\.xlsx$",
                                 default_output_path("cn-north-4")))
        self.assertTrue(re.match(r"ecs_instances_cn-north-4_\d{8}_\d{6}\.csv$",
                                 default_output_path("cn-north-4", "csv")))


class TestExportCsv(unittest.TestCase):
    def _servers(self):
        return [normalize_server(SERVER_1), normalize_server(SERVER_2)]

    def test_csv_header_and_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ecs.csv")
            written = export_csv(self._servers(), path, summary=True,
                                 region="cn-north-4")
            self.assertTrue(os.path.isfile(path))
            self.assertEqual(written, os.path.abspath(path))
            with open(path, encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.reader(fh))
            self.assertEqual(rows[0], REPORT_HEADERS)
            self.assertEqual(rows[1][0], "ecs-web-01")
            self.assertEqual(rows[3][0], "实例数量")
            self.assertEqual(rows[3][1], "2")

    def test_csv_has_utf8_bom(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ecs.csv")
            export_csv(self._servers(), path)
            with open(path, "rb") as fh:
                self.assertTrue(fh.read(3) == b"\xef\xbb\xbf")


class TestMain(unittest.TestCase):
    def _run(self, argv, results):
        fake = FakeRunner(results)
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake), \
                mock.patch("sys.stdout") as stdout, \
                mock.patch("sys.stderr") as stderr:
            rc = main(argv)
        out = "".join(c[0][0] for c in stdout.write.call_args_list)
        err = "".join(c[0][0] for c in stderr.write.call_args_list)
        return rc, out, err

    def _page(self, servers):
        return [(json.dumps({"servers": servers, "count": len(servers)}), "", 0)]

    def _run_capture(self, argv, results):
        # capture real stdout for JSON parsing
        fake = FakeRunner(results)
        import io
        buf = io.StringIO()
        with mock.patch("list_ecs_export_sorted.subprocess.run", fake), \
                mock.patch("sys.stdout", buf):
            rc = main(argv)
        return rc, buf.getvalue()

    def test_xlsx_output_writes_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "ecs.xlsx")
            rc, stdout, err = self._run(
                ["--region", "cn-north-4", "--output", out, "--summary"],
                self._page([SERVER_1, SERVER_2]))
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.isfile(out))
            self.assertIn("ecs.xlsx", stdout)
            self.assertIn("2 个实例", stdout)
            self.assertEqual(err, "")

    def test_csv_output_writes_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "ecs.csv")
            rc, stdout, err = self._run(
                ["--region", "cn-north-4", "--format", "csv",
                 "--sort-by", "name", "--order", "desc", "--output", out],
                self._page([SERVER_1, SERVER_2]))
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.isfile(out))
            self.assertIn("ecs.csv", stdout)
            self.assertIn("sort-by=name order=desc", stdout)

    def test_json_sorted_desc(self):
        rc, out = self._run_capture(
            ["--format", "json", "--sort-by", "created", "--order", "desc"],
            self._page([SERVER_1, SERVER_2]))
        self.assertEqual(rc, 0)
        payload = json.loads(out)
        self.assertEqual([s["name"] for s in payload], ["ecs-web-01", "ecs-db-02"])

    def test_json_sorted_asc(self):
        rc, out = self._run_capture(
            ["--format", "json", "--sort-by", "vcpus", "--order", "asc"],
            self._page([SERVER_1, SERVER_2]))
        self.assertEqual(rc, 0)
        payload = json.loads(out)
        self.assertEqual([s["name"] for s in payload], ["ecs-web-01", "ecs-db-02"])

    def test_json_sorted_by_flattened_key_desc(self):
        # Regression: main() used to sort the *raw* wire records, so the
        # flattened sort keys (vcpus / availability_zone / privateIps / …) were
        # absent and the sort silently became a no-op.
        rc, out = self._run_capture(
            ["--format", "json", "--sort-by", "vcpus", "--order", "desc"],
            self._page([SERVER_1, SERVER_2]))
        self.assertEqual(rc, 0)
        payload = json.loads(out)
        self.assertEqual([s["name"] for s in payload], ["ecs-db-02", "ecs-web-01"])

    def test_json_sorted_by_availability_zone_desc(self):
        rc, out = self._run_capture(
            ["--format", "json", "--sort-by", "可用区", "--order", "desc"],
            self._page([SERVER_1, SERVER_2]))
        self.assertEqual(rc, 0)
        payload = json.loads(out)
        self.assertEqual([s["name"] for s in payload], ["ecs-db-02", "ecs-web-01"])

    def test_table_output(self):
        rc, out, err = self._run(
            ["--format", "table"],
            self._page([SERVER_1, SERVER_2]))
        self.assertEqual(rc, 0)
        self.assertIn("ECS 实例列表", out)
        self.assertIn("ecs-web-01", out)
        self.assertIn("ACTIVE", out)

    def test_unknown_sort_field_errors(self):
        rc, out, err = self._run(["--sort-by", "bogus", "--format", "table"], [])
        self.assertEqual(rc, 1)
        self.assertIn("未知的排序字段", err)

    def test_empty_xlsx_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "empty.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], self._page([]))
            self.assertEqual(rc, 0)
            self.assertIn("无 ECS 实例", stdout)
            self.assertFalse(os.path.exists(out))

    def test_empty_table_message(self):
        rc, out, err = self._run(["--region", "cn-south-1", "--format", "table"],
                                 self._page([]))
        self.assertEqual(rc, 0)
        self.assertIn("无 ECS 实例", out)
        self.assertIn("cn-south-1", out)

    def test_auth_error_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "nope.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [("", "403 Forbidden", 1)])
            self.assertEqual(rc, 1)
            self.assertIn("IAM", err)
            self.assertFalse(os.path.exists(out))

    def test_missing_ak_sk_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "nope.xlsx")
            rc, stdout, err = self._run(
                ["--output", out], [("", "ERROR: no AK configured", 1)])
            self.assertEqual(rc, 1)
            self.assertIn("AK/SK", err)
            self.assertFalse(os.path.exists(out))

    def test_hcloud_missing_falls_back_to_sdk(self):
        with mock.patch("list_ecs_export_sorted.list_servers_via_sdk",
                        return_value=[SERVER_1]):
            rc, out, err = self._run(["--format", "table"], [FileNotFoundError()])
        self.assertEqual(rc, 0)
        self.assertIn("ecs-web-01", out)

    def test_write_intent_operation_blocked(self):
        rc, out, err = self._run(["--operation", "DeleteServers"], [])
        self.assertEqual(rc, 1)
        self.assertIn("写操作意图", err)

    def test_write_flag_argv_blocked(self):
        rc, out, err = self._run(["--create-server"], [])
        self.assertEqual(rc, 1)
        self.assertIn("写操作意图", err)

    def test_status_rebuild_not_blocked(self):
        rc, out, err = self._run(
            ["--format", "table", "--status", "REBUILD"],
            self._page([SERVER_2]))
        self.assertEqual(rc, 0)
        self.assertNotIn("写操作意图", err)

    def test_output_path_with_keyword_not_blocked(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "downloads", "report.xlsx")
            rc, stdout, err = self._run(["--output", out],
                                        self._page([SERVER_1]))
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.isfile(out))


class TestErrorHierarchy(unittest.TestCase):
    def test_subclasses(self):
        for exc in (GuardError, ConfigError, AuthError, ApiError, SortError):
            self.assertTrue(issubclass(exc, EcsExportError))


if __name__ == "__main__":
    unittest.main()